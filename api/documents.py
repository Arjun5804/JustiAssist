
from pathlib import Path
import tempfile
from agents.query_classifier import QueryType
from config import TOP_K_STATUTORY, TOP_K_BAIL
from prompts.templates import build_prompt_with_documents, format_document_citation
from llm_provider import call_llm
from api.query import QueryRequest

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from typing import Optional, List, Dict, Any
import uuid

from core.dependencies import deps
from document_session import session_manager
from services.auth import get_current_user
from services.database import get_db_session, Document, User
from core.storage import storage

def extract_text_from_file(file_path: Path, filename: str) -> str:
    """Extract text from uploaded file"""
    content = ""
    
    if filename.endswith('.txt'):
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
    elif filename.endswith('.pdf'):
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    content += page.extract_text() + "\n"
        except ImportError:
            # Fallback if PyPDF2 not installed
            content = "[PDF extraction requires PyPDF2. Install with: pip install PyPDF2]"
        except Exception as e:
            content = f"[Error extracting PDF: {str(e)}]"
    elif filename.endswith('.docx'):
        try:
            from docx import Document
            doc = Document(str(file_path))
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            content = "\n".join(paragraphs)
        except ImportError:
            content = "[DOCX extraction requires python-docx. Install with: pip install python-docx]"
        except Exception as e:
            content = f"[Error extracting DOCX: {str(e)}]"
    else:
        # Try reading as text
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except:
            content = "[Unable to extract text from this file format]"
    
    return content

router = APIRouter()

@router.post("/upload-document")
async def upload_document(
    file: UploadFile = File(...),
    query: str = Form(...),
    mode: str = Form(default="auto"),
    session_id: str = Form(default=None),
    document_type: str = Form(default="OTHER"),
    current_user: User = Depends(get_current_user)
):
    """
    Upload a document (FIR, charge sheet, case summary) for case-specific analysis.
    
    Documents are:
    - Parsed and chunked
    - Indexed in session-level vector store
    - Clearly marked as NON-STATUTORY evidence
    - Stored persistently in Object Storage for authorized access
    """
    # Validate file
    allowed_extensions = ['.txt', '.pdf', '.doc', '.docx']
    file_ext = Path(file.filename).suffix.lower()
    
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type. Allowed: {', '.join(allowed_extensions)}"
        )
    
    # Size limit (5MB)
    max_size = 5 * 1024 * 1024
    contents = await file.read()
    file_size = len(contents)
    if file_size > max_size:
        raise HTTPException(status_code=400, detail="File too large. Maximum 5MB allowed.")
    
    # Generate ID and Key
    document_id = str(uuid.uuid4())
    safe_filename = file.filename.replace('/', '_').replace('\\', '_')
    object_key = f"documents/{current_user.id}/{document_id}/{safe_filename}"
    
    # Upload to Object Storage
    try:
        await storage.upload(object_key, contents)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to upload document to storage: {str(e)}")
        
    # Persist metadata
    db = get_db_session()
    try:
        if not session_id:
            session_id = str(uuid.uuid4())
            
        doc_meta = Document(
            id=document_id,
            user_id=current_user.id,
            session_id=session_id,
            filename=file.filename,
            object_key=object_key,
            document_type=document_type,
            content_type=file.content_type or "application/octet-stream",
            file_size=file_size
        )
        db.add(doc_meta)
        db.commit()
    except Exception as e:
        db.rollback()
        # Rollback storage upload
        try:
            await storage.delete(object_key)
        except:
            pass
        raise HTTPException(status_code=500, detail=f"Failed to persist document metadata: {str(e)}")
    finally:
        db.close()
        
    # Process document (download to temporary working file as required by Phase 7D)
    with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp:
        tmp_path = Path(tmp.name)
    
    try:
        # Download from object storage for processing
        await storage.download_to_file(object_key, tmp_path)
        document_text = extract_text_from_file(tmp_path, file.filename)
    except Exception as e:
        # If processing fails, DO NOT delete the persistent document
        # We just report the error (or continue with what we have if possible)
        # But here, we can't extract text, so we return a processing error
        raise HTTPException(status_code=422, detail=f"Document stored, but processing failed: {str(e)}")
    finally:
        if tmp_path.exists():
            tmp_path.unlink()  # Clean up temp file
    
    if not document_text or len(document_text) < 10:
        raise HTTPException(status_code=422, detail="Document stored, but could not extract valid text.")
    
    # Get or create in-memory session (runtime cache/FAISS index)
    session = session_manager.get_or_create_session(session_id)
    
    # Add document to session (chunked and indexed)
    uploaded_doc = session.add_document(
        filename=file.filename,
        text=document_text,
        document_type=document_type
    )
    
    processing_info = {
        "steps": [
            f"Uploaded document: {file.filename}",
            f"Document type: {document_type} (NON-STATUTORY)",
            f"Created {len(uploaded_doc.chunks)} chunks",
            f"Session ID: {session.session_id}"
        ]
    }
    
    # Classify query
    if mode == "auto":
        classification = await deps.query_classifier.classify(query)
        query_type = classification.query_type
    elif mode == "bail":
        query_type = QueryType.BAIL_QUERY
    else:
        query_type = QueryType.LEGAL_INFO
    
    processing_info["steps"].append(f"Query type: {query_type.value}")
    
    # Reformulate query
    reformulated = deps.query_reformulator.reformulate(query)
    
    # Search session documents (NON-STATUTORY)
    doc_results = session.search(reformulated.enhanced_query, top_k=3)
    processing_info["steps"].append(f"Retrieved {len(doc_results)} chunks from uploaded documents")
    
    # Search statutory index (AUTHORITATIVE)
    statutory_results = deps.vector_store.search_statutory(
        reformulated.enhanced_query,
        top_k=TOP_K_STATUTORY
    ) if deps.vector_store and deps.vector_store.statutory_index else []
    processing_info["steps"].append(f"Retrieved {len(statutory_results)} statutory provisions")
    
    bail_results = []
    if query_type == QueryType.BAIL_QUERY and deps.vector_store and deps.vector_store.bail_index:
        bail_results = deps.vector_store.search_bail(
            reformulated.enhanced_query,
            top_k=TOP_K_BAIL
        )
        processing_info["steps"].append(f"Retrieved {len(bail_results)} bail precedents")
    
    # Format contexts
    statutory_context = [
        {
            "section_number": r.section_number,
            "law_type": r.law_type,
            "text": r.text,
            "source": r.source_dataset,
            "is_statutory": True,
            "is_authoritative": True
        }
        for r in statutory_results
    ]
    
    # Format document context with NON-STATUTORY markers
    document_context = [
        {
            "filename": r.filename,
            "text": r.text,
            "document_type": r.document_type,
            "is_statutory": False,  # EXPLICIT
            "is_authoritative": False,  # EXPLICIT
            "citation_prefix": r.citation_prefix
        }
        for r in doc_results
    ]
    
    bail_context = [
        {"section_number": r.section_number, "law_type": r.law_type, "text": r.text, "metadata": r.metadata}
        for r in bail_results
    ]
    
    # Build prompt with proper document handling
    if query_type == QueryType.BAIL_QUERY:
        # If document is uploaded, use enhanced document-aware prompt instead of rule-based evaluator
        if document_context:
            # Build enhanced prompt that prioritizes document analysis
            prompt = build_prompt_with_documents(
                query=query,
                statutory_context=statutory_context,
                document_context=document_context,
                bail_context=bail_context,
                is_bail_query=True
            )
            
            answer = await call_llm(prompt)
            confidence_score = 0.85  # High confidence when document is available
            
            # Create assessment from document analysis
            bail_assessment = {
                "likelihood": "See Document Analysis Above",
                "status": "Document-Based Analysis",
                "max_punishment": "See document sections",
                "severity_score": 0,
                "applicable_provisions": ["Based on uploaded document"],
                "explanation": "Analysis based on uploaded court document/FIR."
            }
        else:
            # No documents - use rule-based evaluator
            combined_context = statutory_context
            
            evaluation = await deps.bail_evaluator.evaluate(
                query=query,
                statutory_context=combined_context,
                bail_precedents=bail_context,
                custody_duration_days=None,
                offence_sections=None
            )
            
            answer = evaluation.explanation
            confidence_score = evaluation.confidence_score
            bail_assessment = evaluation.to_dict()
    else:
        # Build prompt with document constraints
        prompt = build_prompt_with_documents(
            query=query,
            statutory_context=statutory_context,
            document_context=document_context,
            bail_context=bail_context,
            is_bail_query=False
        )
        
        answer = await call_llm(prompt)
        confidence_score = 0.7
        bail_assessment = None
    
    processing_info["steps"].append("Generated response with document + statutory context")
    
    # Grounding evaluation (checks both sources)
    all_context = statutory_context + [{"text": d["text"]} for d in document_context]
    grounding_eval = deps.feedback_evaluator.evaluate(answer, all_context, query)
    
    # Format citations - STATUTORY first (authoritative), then documents
    citations = []
    
    # Statutory citations (authoritative)
    for r in statutory_results:
        citations.append({
            "section": r.section_number,
            "law_type": r.law_type,
            "text_preview": r.text[:200] + "..." if len(r.text) > 200 else r.text,
            "source": r.source_dataset,
            "relevance_score": round(r.score, 3),
            "is_authoritative": True,
            "is_statutory": True
        })
    
    # Document citations (non-statutory evidence)
    for r in doc_results:
        citations.append(format_document_citation(
            filename=r.filename,
            text_preview=r.text,
            relevance_score=r.score
        ))
    
    return {
        "query": query,
        "session_id": session.session_id,
        "document_name": file.filename,
        "document_type": document_type,
        "document_chunks": len(uploaded_doc.chunks),
        "document_preview": document_text[:500] + "..." if len(document_text) > 500 else document_text,
        "query_type": query_type.value,
        "answer": answer,
        "citations": citations,
        "confidence_score": round(confidence_score, 2),
        "bail_assessment": bail_assessment,
        "grounding_status": grounding_eval.status.value,
        "processing_info": processing_info,
        "document_notice": "⚠️ Uploaded documents are treated as case evidence, NOT statutory law. Citations clearly distinguish between authoritative legal provisions and user-uploaded evidence."
    }


@router.get("/session/{session_id}/documents")
async def list_session_documents(session_id: str, current_user: User = Depends(get_current_user)):
    """List all documents in a session"""
    db = get_db_session()
    try:
        docs = db.query(Document).filter(Document.session_id == session_id).all()
    finally:
        db.close()
        
    if not docs:
        # Check if it exists in memory only (shouldn't happen with new flow)
        session = session_manager.get_session(session_id)
        if not session:
            return {"session_id": session_id, "documents": [], "total_chunks": 0}
            
    # Verify ownership
    if docs and any(d.user_id != current_user.id for d in docs):
        raise HTTPException(status_code=403, detail="Unauthorized access to session documents")
        
    doc_list = [
        {
            "document_id": d.id,
            "filename": d.filename,
            "document_type": d.document_type,
            "file_size": d.file_size,
            "upload_time": d.created_at.timestamp() if d.created_at else 0
        }
        for d in docs
    ]
    
    # Try to get in-memory session stats
    session = session_manager.get_session(session_id)
    total_chunks = len(session.chunk_metadata) if session else 0
    
    return {
        "session_id": session_id,
        "documents": doc_list,
        "total_chunks": total_chunks
    }


@router.get("/session/{session_id}/document/{document_id}/download")
async def download_document(session_id: str, document_id: str, current_user: User = Depends(get_current_user)):
    """Download a document from object storage"""
    db = get_db_session()
    try:
        doc = db.query(Document).filter(Document.id == document_id, Document.session_id == session_id).first()
        if not doc:
            raise HTTPException(status_code=404, detail="Document not found")
        if doc.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Unauthorized")
        
        object_key = doc.object_key
        content_type = doc.content_type
        filename = doc.filename
    finally:
        db.close()
        
    try:
        if not await storage.exists(object_key):
            raise HTTPException(status_code=404, detail="Document not found in storage")
            
        # Return streaming response
        return StreamingResponse(
            storage.stream(object_key),
            media_type=content_type,
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to download document: {str(e)}")


@router.delete("/session/{session_id}/documents")
async def clear_session_documents(session_id: str, current_user: User = Depends(get_current_user)):
    """Clear all documents from a session"""
    db = get_db_session()
    try:
        docs = db.query(Document).filter(Document.session_id == session_id).all()
        
        if docs and any(d.user_id != current_user.id for d in docs):
            raise HTTPException(status_code=403, detail="Unauthorized access to session documents")
            
        doc_count = len(docs)
        
        for doc in docs:
            try:
                await storage.delete(doc.object_key)
            except Exception as e:
                raise Exception(f"Failed to delete object {doc.object_key} from storage: {str(e)}")
            
            db.delete(doc)
            
        db.commit()
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        db.close()
        
    # Clear in-memory session
    session = session_manager.get_session(session_id)
    if session:
        session.clear()
    
    return {
        "session_id": session_id,
        "message": f"Cleared {doc_count} document(s) from session",
        "documents_cleared": doc_count
    }


@router.delete("/session/{session_id}/document/{document_id}")
async def delete_document(session_id: str, document_id: str, current_user: User = Depends(get_current_user)):
    """Delete a specific document from session"""
    db = get_db_session()
    try:
        doc = db.query(Document).filter(Document.id == document_id, Document.session_id == session_id).first()
        if not doc:
            raise HTTPException(status_code=404, detail="Document not found")
        if doc.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Unauthorized")
            
        object_key = doc.object_key
        
        # Delete from storage
        try:
            await storage.delete(object_key)
        except Exception:
            pass # Best effort
            
        db.delete(doc)
        db.commit()
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        db.close()
        
    # Remove from in-memory session
    session = session_manager.get_session(session_id)
    if session and document_id in session.documents:
        del session.documents[document_id]
        # Rebuild index without deleted document chunks
        session.chunk_metadata = [c for c in session.chunk_metadata if c['document_id'] != document_id]
        
        # Rebuild FAISS index
        if session.chunk_metadata:
            session.index = None  # Reset
            texts = [c['text'] for c in session.chunk_metadata]
            model = session._get_embedding_model()
            embeddings = model.encode(texts, convert_to_numpy=True, normalize_embeddings=True).astype('float32')
            import faiss
            session.index = faiss.IndexFlatIP(session._embedding_dim)
            session.index.add(embeddings)
        else:
            session.index = None
    
    return {
        "session_id": session_id,
        "document_id": document_id,
        "message": "Document deleted successfully",
        "remaining_documents": len(session.documents)
    }


@router.post("/session/{session_id}/query")
async def query_with_session(
    session_id: str,
    request: QueryRequest,
    current_user: User = Depends(get_current_user)
):
    """
    Query using existing session documents.
    Documents from previous uploads in this session are automatically included.
    """
    db = get_db_session()
    try:
        docs = db.query(Document).filter(Document.session_id == session_id).all()
        if docs and any(d.user_id != current_user.id for d in docs):
            raise HTTPException(status_code=403, detail="Unauthorized access to session")
    finally:
        db.close()
        
    session = session_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found or expired")
    
    query = request.query.strip()
    
    # Classify and reformulate
    if request.mode == "auto":
        classification = await deps.query_classifier.classify(query)
        query_type = classification.query_type
    elif request.mode == "bail":
        query_type = QueryType.BAIL_QUERY
    else:
        query_type = QueryType.LEGAL_INFO
    
    reformulated = deps.query_reformulator.reformulate(query)
    
    # Search session documents
    doc_results = session.search(reformulated.enhanced_query, top_k=3)
    
    # Search statutory
    statutory_results = deps.vector_store.hybrid_search_statutory(
        reformulated.enhanced_query,
        top_k=TOP_K_STATUTORY
    ) if deps.vector_store and deps.vector_store.statutory_index else []
    
    bail_results = []
    if query_type == QueryType.BAIL_QUERY and deps.vector_store and deps.vector_store.bail_index:
        bail_results = deps.vector_store.search_bail(
            reformulated.enhanced_query,
            top_k=TOP_K_BAIL
        )
    
    # Format contexts
    statutory_context = [
        {"section_number": r.section_number, "law_type": r.law_type, "text": r.text, "source": r.source_dataset}
        for r in statutory_results
    ]
    
    document_context = [
        {"filename": r.filename, "text": r.text, "document_type": r.document_type}
        for r in doc_results
    ]
    
    # Generate response
    prompt = build_prompt_with_documents(
        query=query,
        statutory_context=statutory_context,
        document_context=document_context,
        bail_context=[{"text": r.text} for r in bail_results] if bail_results else None,
        is_bail_query=(query_type == QueryType.BAIL_QUERY)
    )
    
    answer = await call_llm(prompt)
    
    # Format citations
    citations = []
    for r in statutory_results:
        citations.append({
            "section": r.section_number,
            "law_type": r.law_type,
            "text_preview": r.text[:200] + "...",
            "is_authoritative": True
        })
    for r in doc_results:
        citations.append(format_document_citation(r.filename, r.text, r.score))
    
    return {
        "query": query,
        "session_id": session_id,
        "documents_searched": len(session.documents),
        "query_type": query_type.value,
        "answer": answer,
        "citations": citations,
        "confidence_score": 0.75
    }


